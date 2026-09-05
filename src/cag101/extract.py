"""Text extraction from whatever a submitter uploaded.

Deterministic parsers only — no model calls here. Extraction is cached
separately from scoring so that a rubric change re-scores without re-parsing
(and, more importantly, without re-transcribing video, which is the expensive
step).

Every extractor degrades rather than fails: a format we cannot read produces an
empty result with a reason, which surfaces as a flag on the scorecard instead of
silently vanishing.
"""

from __future__ import annotations

import hashlib
import json
import mimetypes
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path

import fitz  # PyMuPDF

from .config import load_config

# ---------------------------------------------------------------------------
# File classification
# ---------------------------------------------------------------------------
PDF_EXT = {".pdf"}
DOC_EXT = {".docx", ".doc", ".rtf", ".odt"}
PPT_EXT = {".pptx", ".ppt", ".odp"}
SHEET_EXT = {".xlsx", ".xls", ".csv", ".ods"}
VIDEO_EXT = {".mp4", ".mkv", ".mov", ".avi", ".webm", ".wmv", ".m4v"}
AUDIO_EXT = {".mp3", ".wav", ".m4a", ".aac", ".ogg"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tiff"}
TEXT_EXT = {".txt", ".md", ".url", ".weblink"}

FORM_HINTS = ("form", "appendix", "submission", "appendix-i", "appendix i")
DECK_HINTS = ("ppt", "deck", "presentation", "slides")
NOTE_HINTS = ("concept", "note", "writeup", "write-up", "proposal", "brief")
WORKFLOW_HINTS = ("workflow", "flow", "process", "diagram", "architecture")


def classify_file(path: Path) -> str:
    """Coarse role of the file within the submission."""
    ext = path.suffix.lower()
    stem = path.stem.lower()

    if ext in VIDEO_EXT:
        return "video"
    if ext in AUDIO_EXT:
        return "audio"
    if ext in PPT_EXT:
        return "deck"
    if ext in SHEET_EXT:
        return "sheet"
    if ext in IMAGE_EXT:
        return "image"
    if ext in TEXT_EXT:
        return "text"
    if ext in PDF_EXT or ext in DOC_EXT:
        if any(h in stem for h in FORM_HINTS):
            return "form"
        if any(h in stem for h in DECK_HINTS):
            return "deck"
        if any(h in stem for h in WORKFLOW_HINTS):
            return "workflow"
        if any(h in stem for h in NOTE_HINTS):
            return "note"
        return "doc"
    return "other"


def sha256_file(path: Path, chunk: int = 1 << 20) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        while block := fh.read(chunk):
            h.update(block)
    return h.hexdigest()


# ---------------------------------------------------------------------------
# Result type
# ---------------------------------------------------------------------------
@dataclass
class ExtractedDoc:
    relative_path: str
    kind: str
    mime: str
    size_bytes: int
    sha256: str
    text: str = ""
    status: str = "ok"                 # ok | empty | unsupported | error | skipped
    detail: str = ""
    ocr_used: bool = False
    media_derived: bool = False
    page_count: int = 0
    embedded_links: list[str] = field(default_factory=list)

    @property
    def char_count(self) -> int:
        return len(self.text)


URL_RE = re.compile(r"https?://[^\s<>\"')\]]+", re.IGNORECASE)


def _find_links(text: str) -> list[str]:
    seen: list[str] = []
    for match in URL_RE.findall(text or ""):
        cleaned = match.rstrip(".,;:")
        if cleaned not in seen:
            seen.append(cleaned)
    return seen


# ---------------------------------------------------------------------------
# Per-format extractors
# ---------------------------------------------------------------------------
def _extract_pdf(path: Path, doc: ExtractedDoc) -> None:
    cfg = load_config()
    threshold = int(cfg.get("extraction.ocr_char_threshold", 50))
    ocr_enabled = bool(cfg.get("extraction.enable_ocr", False))

    with fitz.open(path) as pdf:
        doc.page_count = len(pdf)
        pages: list[str] = []
        scanned_pages: list[int] = []
        for i, page in enumerate(pdf, start=1):
            text = page.get_text("text") or ""
            if len(text.strip()) < threshold:
                scanned_pages.append(i)
                if ocr_enabled:
                    text = _ocr_page(page) or text
                    doc.ocr_used = True
            pages.append(f"[page {i}]\n{text.strip()}")
        doc.text = "\n\n".join(pages).strip()

    if scanned_pages and not ocr_enabled:
        doc.detail = (
            f"{len(scanned_pages)} of {doc.page_count} page(s) appear to be scanned "
            f"images with little extractable text (pages {scanned_pages[:10]}). "
            "OCR is disabled in config; enable extraction.enable_ocr and install "
            "Tesseract to read them."
        )
        if not doc.text.strip():
            doc.status = "empty"
    elif scanned_pages and ocr_enabled:
        doc.detail = f"OCR applied to {len(scanned_pages)} scanned page(s)."


def _ocr_page(page: "fitz.Page") -> str:
    """OCR a single page. Requires pytesseract and the Tesseract binary."""
    try:
        import io

        import pytesseract
        from PIL import Image
    except ImportError:
        return ""
    try:
        pix = page.get_pixmap(dpi=300)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        return pytesseract.image_to_string(img) or ""
    except Exception:
        return ""


def _extract_docx(path: Path, doc: ExtractedDoc) -> None:
    try:
        import docx
    except ImportError:
        doc.status = "unsupported"
        doc.detail = "python-docx is not installed."
        return

    document = docx.Document(str(path))
    parts: list[str] = [p.text.strip() for p in document.paragraphs if p.text.strip()]

    # Tables carry the substance in this form — Section 1.2 co-submitters and the
    # Section 4.2 impact table are both tables. Preserve cell structure.
    for t_index, table in enumerate(document.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            # Word merges cells by repetition; collapse consecutive duplicates.
            deduped: list[str] = []
            for cell in cells:
                if not deduped or deduped[-1] != cell:
                    deduped.append(cell)
            if any(deduped):
                rows.append(" | ".join(deduped))
        if rows:
            parts.append(f"[table {t_index}]\n" + "\n".join(rows))

    doc.text = "\n\n".join(parts).strip()
    if not doc.text:
        doc.status = "empty"
        doc.detail = "No text found in paragraphs or tables."


def _extract_pptx(path: Path, doc: ExtractedDoc) -> None:
    try:
        from pptx import Presentation
    except ImportError:
        doc.status = "unsupported"
        doc.detail = "python-pptx is not installed."
        return

    prs = Presentation(str(path))
    slides: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        chunks: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                chunks.append(shape.text_frame.text.strip())
            elif getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        chunks.append(" | ".join(cells))
        # Speaker notes often contain the reasoning that the slide omits.
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes:
                chunks.append(f"[speaker notes] {notes}")
        if chunks:
            slides.append(f"[slide {i}]\n" + "\n".join(chunks))
    doc.page_count = len(prs.slides)
    doc.text = "\n\n".join(slides).strip()
    if not doc.text:
        doc.status = "empty"
        doc.detail = "Slides contain no extractable text (likely images only)."


def _extract_sheet(path: Path, doc: ExtractedDoc) -> None:
    if path.suffix.lower() == ".csv":
        doc.text = path.read_text(encoding="utf-8", errors="replace")[:100_000]
        return
    try:
        import openpyxl
    except ImportError:
        doc.status = "unsupported"
        doc.detail = "openpyxl is not installed."
        return
    try:
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    except Exception as exc:
        doc.status = "unsupported"
        doc.detail = f"Could not open workbook: {exc}"
        return
    parts: list[str] = []
    for ws in wb.worksheets:
        rows: list[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"[sheet {ws.title}]\n" + "\n".join(rows[:500]))
    wb.close()
    doc.text = "\n\n".join(parts).strip()


def _extract_text(path: Path, doc: ExtractedDoc) -> None:
    doc.text = path.read_text(encoding="utf-8", errors="replace")[:200_000].strip()
    if not doc.text:
        doc.status = "empty"


def _extract_media(path: Path, doc: ExtractedDoc) -> None:
    """Video/audio. Transcription is opt-in because it needs ffmpeg + Whisper."""
    cfg = load_config()
    if not bool(cfg.get("extraction.enable_video_transcription", False)):
        doc.status = "skipped"
        doc.detail = (
            "Media file present but not transcribed: "
            "extraction.enable_video_transcription is off. Install ffmpeg and "
            "faster-whisper, then enable it in config.yaml. Until then this "
            "submission is scored on its written material only, and is flagged "
            "for human review."
        )
        return

    try:
        from .media import transcribe_media
    except ImportError as exc:
        doc.status = "unsupported"
        doc.detail = f"Transcription unavailable: {exc}"
        return

    try:
        result = transcribe_media(path)
    except Exception as exc:  # transcription must never kill a run
        doc.status = "error"
        doc.detail = f"Transcription failed: {exc}"
        return

    limit = int(cfg.get("extraction.max_transcript_chars", 40_000))
    doc.text = result.text[:limit]
    doc.media_derived = True
    doc.detail = result.detail


def _extract_image(path: Path, doc: ExtractedDoc) -> None:
    """Images are not read here. A diagram or screenshot is described by the
    vision model during digest building, where it can be given context."""
    doc.status = "skipped"
    doc.detail = "Image recorded; described during digest building by the vision model."


UNSUPPORTED_LEGACY = {
    ".doc": "Legacy .doc format. Convert to .docx or PDF to make it readable.",
    ".ppt": "Legacy .ppt format. Convert to .pptx or PDF to make it readable.",
    ".rtf": "RTF is not parsed. Convert to .docx or PDF.",
    ".odt": "OpenDocument text is not parsed. Convert to .docx or PDF.",
    ".odp": "OpenDocument presentation is not parsed. Convert to .pptx or PDF.",
}


def extract_file(path: Path, submission_root: Path) -> ExtractedDoc:
    """Extract one file. Never raises — failures come back as status + detail."""
    rel = str(path.relative_to(submission_root)).replace("\\", "/")
    mime = mimetypes.guess_type(path.name)[0] or ""
    try:
        size = path.stat().st_size
        digest = sha256_file(path)
    except OSError as exc:
        return ExtractedDoc(
            relative_path=rel, kind="other", mime=mime, size_bytes=0, sha256="",
            status="error", detail=f"Could not read file: {exc}",
        )

    doc = ExtractedDoc(
        relative_path=rel,
        kind=classify_file(path),
        mime=mime,
        size_bytes=size,
        sha256=digest,
    )

    ext = path.suffix.lower()
    try:
        if ext in UNSUPPORTED_LEGACY:
            doc.status = "unsupported"
            doc.detail = UNSUPPORTED_LEGACY[ext]
        elif ext in PDF_EXT:
            _extract_pdf(path, doc)
        elif ext == ".docx":
            _extract_docx(path, doc)
        elif ext == ".pptx":
            _extract_pptx(path, doc)
        elif ext in SHEET_EXT:
            _extract_sheet(path, doc)
        elif ext in TEXT_EXT:
            _extract_text(path, doc)
        elif ext in VIDEO_EXT or ext in AUDIO_EXT:
            _extract_media(path, doc)
        elif ext in IMAGE_EXT:
            _extract_image(path, doc)
        else:
            doc.status = "unsupported"
            doc.detail = f"No extractor for {ext or 'files without an extension'}."
    except Exception as exc:
        doc.status = "error"
        doc.detail = f"{type(exc).__name__}: {exc}"

    doc.embedded_links = _find_links(doc.text)
    return doc


# ---------------------------------------------------------------------------
# Cache
# ---------------------------------------------------------------------------
def _cache_path(submission_ref: str) -> Path:
    return load_config().path("paths.cache") / f"{submission_ref}_extraction.json"


def load_cached_extraction(submission_ref: str, expected_hash: str) -> list[ExtractedDoc] | None:
    """Return cached extraction if the submission folder has not changed."""
    path = _cache_path(submission_ref)
    if not path.exists():
        return None
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return None
    if payload.get("content_hash") != expected_hash:
        return None
    return [ExtractedDoc(**d) for d in payload.get("documents", [])]


def save_extraction(submission_ref: str, content_hash: str, docs: list[ExtractedDoc]) -> Path:
    path = _cache_path(submission_ref)
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "submission_ref": submission_ref,
        "content_hash": content_hash,
        "documents": [asdict(d) for d in docs],
    }
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return path
