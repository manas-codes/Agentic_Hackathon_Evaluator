"""Excel export — the artefact a Filtering Committee actually works from.

Three sheets, no more:

  Summary     what was scored, against which rubric, and how the marks fell
  Scorecards  one row per submission — the portal's table, plus only the
              columns an evaluator needs when working away from the screen
  Evidence    long format, one row per sub-score with its quote and source, so
              any figure can be verified without opening the portal

The Evidence sheet is the one that makes this defensible. A scorecard without
traceable evidence is an opinion; with it, a committee can audit any figure.

Nothing here names the tooling that produced the assessment. That belongs in
the database and the audit log, not on paper going round a committee, where it
invites argument about the instrument instead of the submission.

Submitter names and offices appear only when the person downloading the
workbook is entitled to see them in the portal — the file must not be a way
around the role that governs the screen.
"""

from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .config import ensure_directories, load_config
from .db import init_db, session_scope
from .reporting import (
    FORMATION_LABELS,
    PUBLIC_EVALUATOR_LABEL,
    STAGE_LABELS,
    THEME_LABELS,
    ScorecardView,
    all_scorecards,
    build_overview,
    public_flags,
    rank_scorecards,
)
from .rubric import load_rubric

# Timestamps are stored in UTC; the workbook is read in India, so dates are
# rendered in IST to match the screens.
_IST = timezone(timedelta(hours=5, minutes=30))


def _ist_date(value: datetime | None) -> str:
    if value is None:
        return ""
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(_IST).strftime("%d-%m-%Y")


# --- styling ---------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
SUBTLE = Font(size=9, color="666666")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
CENTRE = Alignment(horizontal="center", vertical="center")
THIN = Side(style="thin", color="D9D9D9")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")
REVIEW_FILL = PatternFill("solid", fgColor="FCE4D6")


def _style_header(ws: Worksheet, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BOX
    ws.row_dimensions[row].height = 34


def _widths(ws: Worksheet, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


# ---------------------------------------------------------------------------
def export_workbook(
    out_path: Path | None = None, include_identity: bool = False
) -> Path:
    """Build the workbook.

    `include_identity` must carry the downloader's actual entitlement. It
    defaults to False so that a caller who forgets to pass it produces the
    safe file, not the one that leaks.
    """
    init_db()
    ensure_directories()
    cfg = load_config()
    rubric = load_rubric()

    with session_scope() as session:
        cards = all_scorecards(
            session, rubric, include_identity=include_identity, include_evidence=True
        )
        rank_scorecards(cards)
        overview = build_overview(session, cards, rubric)

    if out_path is None:
        stamp = datetime.now().strftime("%Y%m%d-%H%M")
        out_path = cfg.path("paths.exports") / f"CAG101-Scorecards-{stamp}.xlsx"
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    _sheet_summary(wb.active, overview, rubric)
    _sheet_scorecards(wb.create_sheet("Scorecards"), cards, rubric, include_identity)
    _sheet_evidence(wb.create_sheet("Evidence"), cards)

    wb.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
def _sheet_summary(ws: Worksheet, ov, rubric) -> None:
    ws.title = "Summary"
    _widths(ws, [42, 22, 22, 22])

    ws["A1"] = "CAG 101 Innovation Ideas Initiative — Evaluation Summary"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (
        "Indicative pre-assessment to assist the Filtering Committees. "
        "Not a decision."
    )
    ws["A2"].font = SUBTLE

    row = 4
    for label, value in [
        ("Generated", datetime.now().strftime("%d-%m-%Y %H:%M")),
        ("Rubric", f"{ov.rubric_name} ({ov.rubric_slug})"),
        ("Rubric content hash", ov.rubric_hash),
        ("Total submissions", ov.total_submissions),
        ("Evaluated", ov.evaluated),
        ("Pending evaluation", ov.pending),
        ("Failed / needs attention", ov.failed),
        ("Flagged for human review", ov.needs_review),
        ("With a human override", ov.overridden),
        ("Indicative NAeG candidates", ov.naeg_candidates),
    ]:
        ws.cell(row=row, column=1, value=label).font = BOLD
        ws.cell(row=row, column=2, value=value)
        row += 1

    if ov.evaluated:
        row += 1
        ws.cell(row=row, column=1, value="Score distribution").font = TITLE_FONT
        row += 1
        for label, value in [
            ("Mean", ov.mean_score),
            ("Median", ov.median_score),
            ("Lowest", ov.min_score),
            ("Highest", ov.max_score),
        ]:
            ws.cell(row=row, column=1, value=label).font = BOLD
            ws.cell(row=row, column=2, value=value)
            row += 1

        row += 1
        ws.cell(row=row, column=1, value="Band").font = HEADER_FONT
        ws.cell(row=row, column=2, value="Submissions").font = HEADER_FONT
        _style_header(ws, row, 2)
        row += 1
        for label, count in ov.histogram:
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=count)
            row += 1

        row += 1
        ws.cell(row=row, column=1, value="Mean score by criterion").font = TITLE_FONT
        row += 1
        ws.cell(row=row, column=1, value="Criterion")
        ws.cell(row=row, column=2, value="Mean")
        ws.cell(row=row, column=3, value="Maximum")
        _style_header(ws, row, 3)
        row += 1
        for name, mean, maximum in ov.criterion_means:
            ws.cell(row=row, column=1, value=name)
            ws.cell(row=row, column=2, value=mean)
            ws.cell(row=row, column=3, value=maximum)
            row += 1

    for heading, data in [
        ("Submissions by thematic area", ov.by_theme),
        ("Submissions by formation stream", ov.by_formation),
        ("Submissions by development stage", ov.by_stage),
        ("Submissions by completeness", ov.by_completeness),
        ("Submissions by evidence strength", ov.by_evidence_strength),
    ]:
        if not data:
            continue
        row += 1
        ws.cell(row=row, column=1, value=heading).font = TITLE_FONT
        row += 1
        for key, count in data.items():
            ws.cell(row=row, column=1, value=key)
            ws.cell(row=row, column=2, value=count)
            row += 1

    if ov.last_run:
        row += 1
        ws.cell(row=row, column=1, value="Last scoring run").font = TITLE_FONT
        row += 1
        for label, key in [
            ("Run id", "run_id"),
            ("Submissions attempted", "total"),
            ("Completed", "done"),
            ("Failed", "failed"),
            ("State", "state"),
        ]:
            ws.cell(row=row, column=1, value=label).font = BOLD
            ws.cell(row=row, column=2, value=ov.last_run.get(key))
            row += 1

    ws.freeze_panes = "A4"


# ---------------------------------------------------------------------------
def _sheet_scorecards(
    ws: Worksheet, cards: list[ScorecardView], rubric, include_identity: bool = False
) -> None:
    """One row per submission — the portal's table, in a spreadsheet.

    The columns are the ones on screen, plus only what an evaluator genuinely
    needs when working away from the portal: where sign-off stands, why a
    submission is flagged, and the remark that explains the total. The twelve
    sub-scores are deliberately not here — they belong with their quotes on the
    Evidence sheet, and a hundred-column grid is not something a committee
    reads.
    """
    headers = ["Ref", "Rank", "Submission"]
    if include_identity:
        headers.append("Office")
    headers += ["Formation stream", "Thematic area", "Development stage"]
    for criterion in rubric.criteria:
        headers.append(f"{criterion.name} (/{criterion.weight})")
    headers += [
        "TOTAL (/100)",
        "Amended from",
        "Committee decision", "Approved by", "Approved on",
        "Needs review", "Why flagged",
        "Remarks — why this total",
        "Status", "Files",
    ]

    ws.append(headers)
    _style_header(ws, 1, len(headers))

    for card in cards:
        row: list = [
            card.ref,
            getattr(card, "rank_overall", None),
            card.title or card.folder_name,
        ]
        if include_identity:
            row.append(card.office)
        row += [
            FORMATION_LABELS.get(card.formation_category, card.formation_category),
            THEME_LABELS.get(card.thematic_area, card.thematic_area),
            STAGE_LABELS.get(card.development_stage, card.development_stage),
        ]
        for criterion in rubric.criteria:
            view = card.criterion(criterion.key)
            row.append(view.effective_score if view and card.evaluated else None)

        row += [
            card.total if card.evaluated else None,
            # blank unless an evaluator changed a mark, so the column is silent
            # on the great majority of rows rather than repeating the total
            card.machine_total if card.is_overridden else None,
            card.review_decision,
            card.review_reviewer,
            _ist_date(card.review_decided_at),
            "YES" if card.needs_review else "",
            next((r["detail"] for r in public_flags(card)
                  if r["key"] == "human review required"), ""),
            card.remarks,
            card.status.replace("_", " "),
            len(card.files),
        ]
        ws.append(row)

        excel_row = ws.max_row
        if card.needs_review:
            for col in range(1, len(headers) + 1):
                ws.cell(row=excel_row, column=col).fill = REVIEW_FILL
        elif card.is_overridden:
            for col in range(1, len(headers) + 1):
                ws.cell(row=excel_row, column=col).fill = FLAG_FILL

    widths = [13, 7, 52]
    if include_identity:
        widths.append(34)
    widths += [20, 30, 20]
    widths += [16] * len(rubric.criteria)
    widths += [14, 15, 18, 20, 14, 13, 46, 70, 18, 7]
    _widths(ws, widths[: len(headers)])

    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 2)}"

    total_col = headers.index("TOTAL (/100)") + 1
    for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row_cells:
            cell.alignment = TOP
        row_cells[2].alignment = WRAP                  # submission title
        ws.cell(row=row_cells[0].row, column=total_col).font = Font(bold=True, size=12)


# ---------------------------------------------------------------------------
def _sheet_evidence(ws: Worksheet, cards: list[ScorecardView]) -> None:
    headers = [
        "Ref", "Title", "Criterion", "Sub-criterion", "Score", "Max", "Band",
        "Reasoning", "Evidence quote", "Source file", "Locator",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    for card in cards:
        if not card.evaluated:
            continue
        for criterion in card.criteria:
            for sub in criterion.sub_scores:
                if sub.evidence:
                    for e in sub.evidence:
                        ws.append([
                            card.ref, card.title or card.folder_name,
                            criterion.name, sub.sub_name, sub.score, sub.max_score,
                            sub.band, sub.rationale,
                            e.get("quote", ""), e.get("source_file", ""),
                            e.get("locator", ""),
                        ])
                else:
                    # An unevidenced sub-score must appear here too, marked as
                    # such — its absence from this sheet is the finding.
                    ws.append([
                        card.ref, card.title or card.folder_name,
                        criterion.name, sub.sub_name, sub.score, sub.max_score,
                        sub.band, sub.rationale,
                        "No supporting quote found", "", "",
                    ])
                    ws.cell(row=ws.max_row, column=9).fill = REVIEW_FILL

    _widths(ws, [13, 38, 28, 34, 8, 7, 13, 60, 70, 30, 16])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 2)}"
    for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row_cells:
            cell.alignment = TOP
        row_cells[7].alignment = WRAP    # reasoning
        row_cells[8].alignment = WRAP    # quote
